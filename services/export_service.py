"""Inspection export: CSV, Excel (openpyxl) and PDF (reportlab).

All exporters take detached ORM ``Inspection`` rows (details eagerly loaded,
as returned by ``InspectionRepository.search``) so the Database Viewer exports
exactly what it displays.
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from core.database import Inspection
from core.logging import get_logger
from core.utilities.enums import LogSource
from core.utilities.exceptions import ExportError

logger = get_logger(LogSource.SYSTEM)

CAMERA_INDEXES = (1, 2, 3, 4)


def _headers() -> list[str]:
    headers = ["ID", "Date", "Time", "Machine No", "Serial No", "Overall"]
    for index in CAMERA_INDEXES:
        headers += [f"Cam{index} X (mm)", f"Cam{index} Y (mm)", f"Cam{index} Result"]
    headers += ["Cycle (ms)", "Detect (ms)", "Operator", "Shift"]
    return headers


def _row(inspection: Inspection) -> list:
    details = {detail.camera_index: detail for detail in inspection.details}
    row: list = [
        inspection.id,
        inspection.created_at.strftime("%Y-%m-%d"),
        inspection.created_at.strftime("%H:%M:%S"),
        inspection.machine_number,
        inspection.serial_number,
        inspection.overall_result,
    ]
    for index in CAMERA_INDEXES:
        detail = details.get(index)
        if detail is None:
            row += ["", "", ""]
        elif not detail.hole_found:
            row += ["-", "-", detail.result]
        else:
            row += [f"{detail.x_mm:.2f}", f"{detail.y_mm:.2f}", detail.result]
    row += [
        f"{inspection.plc_cycle_time_ms:.0f}",
        f"{inspection.detection_time_ms:.0f}",
        inspection.operator,
        inspection.shift,
    ]
    return row


class ExportService:
    """Stateless exporters; every method raises ExportError on failure."""

    def export_csv(self, path: str | Path, inspections: list[Inspection]) -> Path:
        path = Path(path)
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            # utf-8-sig so Excel opens the file with correct encoding
            with path.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.writer(handle)
                writer.writerow(_headers())
                writer.writerows(_row(inspection) for inspection in inspections)
        except OSError as exc:
            raise ExportError(f"CSV export failed: {exc}") from exc
        logger.info("Exported %d inspections to %s", len(inspections), path)
        return path

    def export_excel(self, path: str | Path, inspections: list[Inspection]) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise ExportError("openpyxl is not installed") from exc

        path = Path(path)
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Inspections"

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1F3B57")
            headers = _headers()
            sheet.append(headers)
            for column, _ in enumerate(headers, start=1):
                cell = sheet.cell(row=1, column=column)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
                sheet.column_dimensions[get_column_letter(column)].width = (
                    max(11, len(headers[column - 1]) + 2)
                )
            sheet.freeze_panes = "A2"

            for inspection in inspections:
                sheet.append(_row(inspection))
            workbook.save(path)
        except OSError as exc:
            raise ExportError(f"Excel export failed: {exc}") from exc
        logger.info("Exported %d inspections to %s", len(inspections), path)
        return path

    def export_pdf(
        self,
        path: str | Path,
        inspections: list[Inspection],
        title: str = "Inspection Report",
    ) -> Path:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError as exc:
            raise ExportError("reportlab is not installed") from exc

        path = Path(path)
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            document = SimpleDocTemplate(
                str(path),
                pagesize=landscape(A4),
                leftMargin=10 * mm,
                rightMargin=10 * mm,
                topMargin=12 * mm,
                bottomMargin=12 * mm,
                title=title,
            )
            styles = getSampleStyleSheet()
            generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            story = [
                Paragraph(title, styles["Title"]),
                Paragraph(
                    f"Generated {generated} — {len(inspections)} inspections",
                    styles["Normal"],
                ),
                Spacer(1, 6 * mm),
            ]

            data = [_headers()] + [_row(inspection) for inspection in inspections]
            table = Table(data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3B57")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                         [colors.white, colors.HexColor("#EDF1F5")]),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            story.append(table)
            document.build(story)
        except OSError as exc:
            raise ExportError(f"PDF export failed: {exc}") from exc
        logger.info("Exported %d inspections to %s", len(inspections), path)
        return path
